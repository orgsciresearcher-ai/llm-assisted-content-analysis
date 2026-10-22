"""Create blind coding sheets (one Excel file per coder) from a sample drawn with draw_samples.py.

The sheets never show model codes, confidence, rationales or strata.

  segments      confidence-calibration test: each row is one segment; coders apply the full codebook
  adjudication  theory-label test: each row is one document-code cell; coders judge whether the text
                supports the code, without knowing which prompt assigned it
  documents     recall check and document-level coding: each row is one document; coders read it in full
                and apply the full codebook

Examples
  python tools/make_blind_sheets.py segments --sample case/samples/calibration_test_sample.csv \
      --segments case/data/coded_segments.csv --index case/data/ads_index.csv \
      --codebook codebook/case/codebook_A_original.csv --corpus PATH/TO/full_texts.csv --corpus-id-col ad_id \
      --corpus-text-col full_text --coders coder1,coder2 --out-dir sheets
  python tools/make_blind_sheets.py adjudication --cells samples/adjudication_cells.csv --index case/data/ads_index.csv \
      --codebook codebook/case/codebook_A_original.csv --coders coder1,coder2 --out-dir sheets
  python tools/make_blind_sheets.py documents --docs case/samples/recall_check_uncoded_ads.csv \
      --index case/data/ads_index.csv --codebook codebook/case/codebook_A_original.csv --out-dir sheets

The optional full-text corpus (kept locally, never shared) lets the script add the text just before each
segment and a guess of its section heading. Coders can always open the link to the full document.
"""
import argparse
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEAD = PatternFill("solid", start_color="3979A8")


def guess_heading(full_text, segment, context_chars=300):
    """(heading guess, preceding text). The guess is the nearest short line above the segment that does
    not end like a sentence; the preceding text lets coders check it. Empty strings if not found."""
    if not isinstance(full_text, str) or not isinstance(segment, str):
        return "", ""
    low = full_text.lower()
    pos = -1
    words = re.findall(r"\w+", segment.lower())
    for n in (12, 8, 5):
        if len(words) >= n:
            pat = r"\W+".join(re.escape(w) for w in words[:n])
            m = re.search(pat, low)
            if m:
                pos = m.start()
                break
    if pos < 0:
        return "", ""
    before = full_text[max(0, pos - context_chars):pos].strip()
    for line in reversed(full_text[:pos].splitlines()[:-1] or []):
        ln = line.strip().strip(":").strip()
        if 2 <= len(ln) <= 80 and not ln.endswith((".", ";", ",")) and not re.match(r"^\d{4}-\d{2}-\d{2}", ln):
            return ln, before
    return "", before


def new_book():
    wb = openpyxl.Workbook()
    wb.properties.creator = "anonymous"
    wb.properties.lastModifiedBy = "anonymous"
    return wb


def header(ws, cols, widths):
    ws.append(cols)
    for j, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=j)
        c.fill, c.font = HEAD, Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def instructions(wb, lines):
    ws = wb.create_sheet("Instructions")
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 110


def cmd_segments(a):
    sample = pd.read_csv(a.sample)
    seg = pd.read_csv(a.segments).set_index("seg_id")
    idx = pd.read_csv(a.index).set_index("ad_id") if a.index else None
    codes = pd.read_csv(a.codebook)["code"].astype(str).tolist()
    texts = {}
    if a.corpus:
        c = pd.read_csv(a.corpus)
        texts = dict(zip(c[a.corpus_id_col].astype(str), c[a.corpus_text_col]))
    rng = np.random.default_rng(a.seed)
    order = rng.permutation(len(sample))
    rows = []
    for item, k in enumerate(order, 1):
        sid = sample.iloc[k]["seg_id"]
        ad = str(seg.at[sid, "ad_id"])
        text = seg.at[sid, "text_segment"]
        heading, before = guess_heading(texts.get(ad), text)
        rows.append([item, sid, idx.at[ad, "url"] if idx is not None and ad in idx.index else "",
                     heading, before, text])
    key = pd.DataFrame({"item": [r[0] for r in rows], "seg_id": [r[1] for r in rows]})
    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    key.to_csv(out / "segments_item_key.csv", index=False)
    for coder in a.coders.split(","):
        wb = new_book()
        ws = wb.active
        ws.title = "Coding"
        header(ws, ["item", "seg_id", "document_link", "section_heading_guess", "text_before_segment", "segment_text"]
               + codes + ["notes"], [6, 14, 30, 24, 40, 70] + [8] * len(codes) + [30])
        for r in rows:
            ws.append(r + [None] * len(codes) + [None])
            for col in (5, 6):
                ws.cell(row=ws.max_row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
        ws.add_data_validation(dv)
        first = openpyxl.utils.get_column_letter(7)
        last = openpyxl.utils.get_column_letter(6 + len(codes))
        dv.add(f"{first}2:{last}{len(rows) + 1}")
        instructions(wb, [
            f"Blind coding sheet for {coder} (confidence-calibration test)",
            "Code every segment independently with the full codebook. Enter 1 in each code column that applies.",
            "A segment may receive several codes. Leave all code columns empty if no code applies.",
            "Do not look up the model output, its confidence, or the stratum of a segment.",
            "Do not discuss segments with the other coder until both sheets are complete.",
        ])
        wb.save(out / f"blind_segments_{coder}.xlsx")
    print(f"{len(rows)} segments -> {out} ({a.coders})")


def cmd_adjudication(a):
    cells = pd.read_csv(a.cells)
    idx = pd.read_csv(a.index).set_index("ad_id") if a.index else None
    cb = pd.read_csv(a.codebook).set_index("code")
    rng = np.random.default_rng(a.seed)
    order = rng.permutation(len(cells))
    rows = []
    for item, k in enumerate(order, 1):
        r = cells.iloc[k]
        doc, code = str(r["doc_id"]), str(r["code"])
        definition = f"{cb.at[code, 'definition']} Code when: {cb.at[code, 'code_when']}" if code in cb.index else ""
        rows.append([item, doc, idx.at[doc, "url"] if idx is not None and doc in idx.index else "", code, definition])
    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    pd.DataFrame({"item": [r[0] for r in rows], "doc_id": [r[1] for r in rows], "code": [r[3] for r in rows]}
                 ).to_csv(out / "adjudication_item_key.csv", index=False)
    for coder in a.coders.split(","):
        wb = new_book()
        ws = wb.active
        ws.title = "Adjudication"
        header(ws, ["item", "doc_id", "document_link", "code", "code_definition",
                    "supported_by_text", "evidence_quote", "notes"], [6, 12, 30, 9, 60, 16, 50, 30])
        for r in rows:
            ws.append(r + [None, None, None])
            ws.cell(row=ws.max_row, column=5).alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"yes,no,unclear"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"F2:F{len(rows) + 1}")
        instructions(wb, [
            f"Blind adjudication sheet for {coder} (theory-label test)",
            "Each row is one document-code cell. Decide from the document text alone whether the code applies.",
            "Quote the evidence if it does. The sheet does not show which prompt assigned the code.",
        ])
        wb.save(out / f"blind_adjudication_{coder}.xlsx")
    print(f"{len(rows)} cells -> {out} ({a.coders})")


def cmd_documents(a):
    docs = pd.read_csv(a.docs)
    idx = pd.read_csv(a.index).set_index(a.id_col) if a.index else None
    codes = pd.read_csv(a.codebook)["code"].astype(str).tolist()
    rng = np.random.default_rng(a.seed)
    order = rng.permutation(len(docs))
    rows = []
    for item, k in enumerate(order, 1):
        doc = str(docs.iloc[k][a.id_col])
        rows.append([item, doc, idx.at[doc, "url"] if idx is not None and doc in idx.index else ""])
    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    pd.DataFrame({"item": [r[0] for r in rows], "doc_id": [r[1] for r in rows]}
                 ).to_csv(out / "documents_item_key.csv", index=False)
    for coder in a.coders.split(","):
        wb = new_book()
        ws = wb.active
        ws.title = "Coding"
        header(ws, ["item", "doc_id", "document_link"] + codes + ["evidence_quotes", "notes"],
               [6, 12, 40] + [8] * len(codes) + [60, 30])
        for r in rows:
            ws.append(r + [None] * len(codes) + [None, None])
        dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
        ws.add_data_validation(dv)
        first = openpyxl.utils.get_column_letter(4)
        last = openpyxl.utils.get_column_letter(3 + len(codes))
        dv.add(f"{first}2:{last}{len(rows) + 1}")
        instructions(wb, [
            f"Blind coding sheet for {coder} (document level)",
            "Read each document in full and apply the full codebook, including its eligibility rules.",
            "Enter 1 in each code column that applies anywhere in the document, and quote the passage in evidence_quotes.",
            "Leave all code columns empty if no code applies.",
            "Do not look up the model output. Do not discuss documents with the other coder until both sheets are complete.",
        ])
        wb.save(out / f"blind_documents_{coder}.xlsx")
    print(f"{len(rows)} documents -> {out} ({a.coders})")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("segments")
    s.add_argument("--sample", required=True)
    s.add_argument("--segments", required=True)
    s.add_argument("--index")
    s.add_argument("--codebook", required=True)
    s.add_argument("--corpus")
    s.add_argument("--corpus-id-col", default="ad_id")
    s.add_argument("--corpus-text-col", default="full_text")
    s.add_argument("--coders", default="coder1,coder2")
    s.add_argument("--seed", type=int, default=20261101)
    s.add_argument("--out-dir", required=True)
    s.set_defaults(func=cmd_segments)
    d = sub.add_parser("adjudication")
    d.add_argument("--cells", required=True)
    d.add_argument("--index")
    d.add_argument("--codebook", required=True)
    d.add_argument("--coders", default="coder1,coder2")
    d.add_argument("--seed", type=int, default=20261101)
    d.add_argument("--out-dir", required=True)
    d.set_defaults(func=cmd_adjudication)
    e = sub.add_parser("documents")
    e.add_argument("--docs", required=True, help="CSV with one row per document to code")
    e.add_argument("--id-col", default="ad_id")
    e.add_argument("--index")
    e.add_argument("--codebook", required=True)
    e.add_argument("--coders", default="coder1,coder2")
    e.add_argument("--seed", type=int, default=20261101)
    e.add_argument("--out-dir", required=True)
    e.set_defaults(func=cmd_documents)
    a = p.parse_args()
    a.func(a)


if __name__ == "__main__":
    main()
